"""Build a read-only metadata and schema snapshot of the project Dropbox corpus.

The script never modifies source files. It excludes Graphify-generated sidecars in
``graphify-out`` so that coverage statistics describe the user-owned corpus.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import struct
import subprocess
import zipfile
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def dbf_profile(path: Path) -> dict:
    """Read DBF header metadata without optional GIS dependencies."""
    with path.open("rb") as stream:
        header = stream.read(32)
        if len(header) != 32:
            raise ValueError("DBF header is incomplete")
        records = struct.unpack("<I", header[4:8])[0]
        header_length = struct.unpack("<H", header[8:10])[0]
        record_length = struct.unpack("<H", header[10:12])[0]
        fields = []
        while stream.tell() < header_length - 1:
            descriptor = stream.read(32)
            if not descriptor or descriptor[0] == 0x0D:
                break
            name = descriptor[:11].split(b"\x00", 1)[0].decode("latin-1")
            fields.append(
                {
                    "name": name,
                    "type": chr(descriptor[11]),
                    "length": int(descriptor[16]),
                    "decimals": int(descriptor[17]),
                }
            )
    return {
        "records": records,
        "record_length": record_length,
        "fields": fields,
    }


def shapefile_profile(path: Path) -> dict:
    """Read the ESRI shapefile header and count records dependency-free."""
    with path.open("rb") as stream:
        header = stream.read(100)
        if len(header) != 100 or struct.unpack(">I", header[:4])[0] != 9994:
            raise ValueError("Invalid shapefile header")
        file_length_bytes = struct.unpack(">I", header[24:28])[0] * 2
        version, shape_type = struct.unpack("<II", header[28:36])
        bbox = list(struct.unpack("<4d", header[36:68]))
        records = 0
        while stream.tell() + 8 <= file_length_bytes:
            record_header = stream.read(8)
            if len(record_header) < 8:
                break
            content_length = struct.unpack(">I", record_header[4:8])[0] * 2
            stream.seek(content_length, os.SEEK_CUR)
            records += 1
    result = {
        "records": records,
        "version": version,
        "geometry_type": shape_type,
        "bbox": bbox,
    }
    sidecar = path.with_suffix(".dbf")
    if sidecar.exists():
        result["attribute_table"] = dbf_profile(sidecar)
    return result


def legacy_xls_profile(path: Path) -> list[dict]:
    """Profile a legacy .xls workbook through the installed read-only ACE provider."""
    escaped = str(path).replace("'", "''")
    script = rf"""
$path='{escaped}'
$cs="Provider=Microsoft.ACE.OLEDB.16.0;Data Source=$path;Extended Properties='Excel 8.0;HDR=YES;IMEX=1';Mode=Read;"
$cn=New-Object System.Data.OleDb.OleDbConnection($cs)
try {{
  $cn.Open()
  $tables=$cn.GetSchema('Tables') | Where-Object {{ $_.TABLE_TYPE -eq 'TABLE' -and $_.TABLE_NAME -notlike '*FilterDatabase*' }}
  $out=@()
  foreach($t in $tables) {{
    $name=$t.TABLE_NAME
    $q=$cn.CreateCommand(); $q.CommandText="SELECT COUNT(*) FROM [$name]"
    $count=[int]$q.ExecuteScalar()
    $q2=$cn.CreateCommand(); $q2.CommandText="SELECT TOP 1 * FROM [$name]"
    $reader=$q2.ExecuteReader(); $schema=$reader.GetSchemaTable()
    $cols=@($schema.Rows | Sort-Object ColumnOrdinal | ForEach-Object {{ $_.ColumnName }})
    $reader.Close()
    $out += [pscustomobject]@{{sheet=$name; rows=$count; columns=$cols}}
  }}
  @($out) | ConvertTo-Json -Depth 5 -Compress
}} finally {{ if($cn.State -eq 'Open'){{$cn.Close()}} }}
"""
    completed = subprocess.run(
        ["powershell.exe", "-NoProfile", "-Command", script],
        check=True,
        capture_output=True,
        text=True,
    )
    parsed = json.loads(completed.stdout)
    return parsed if isinstance(parsed, list) else [parsed]


def basic_profile(path: Path) -> dict:
    suffix = path.suffix.lower()
    profile: dict = {"status": "inventoried", "format": suffix or "[none]"}
    try:
        if suffix == ".dta":
            import pandas as pd

            data = pd.read_stata(path, convert_categoricals=False)
            profile.update(
                status="schema_profiled",
                rows=int(data.shape[0]),
                columns=int(data.shape[1]),
                variable_names=list(map(str, data.columns)),
                dtypes={str(name): str(dtype) for name, dtype in data.dtypes.items()},
            )
        elif suffix == ".csv":
            import pandas as pd

            data = pd.read_csv(path, low_memory=False)
            profile.update(
                status="schema_profiled",
                rows=int(data.shape[0]),
                columns=int(data.shape[1]),
                variable_names=list(map(str, data.columns)),
                dtypes={str(name): str(dtype) for name, dtype in data.dtypes.items()},
            )
        elif suffix == ".xlsx":
            import pandas as pd

            book = pd.ExcelFile(path)
            sheets = []
            for sheet_name in book.sheet_names:
                frame = pd.read_excel(book, sheet_name=sheet_name)
                sheets.append(
                    {
                        "sheet": str(sheet_name),
                        "rows": int(frame.shape[0]),
                        "columns": int(frame.shape[1]),
                        "variable_names": list(map(str, frame.columns)),
                    }
                )
            profile.update(status="schema_profiled", sheets=sheets)
        elif suffix == ".xls":
            profile.update(status="schema_profiled", sheets=legacy_xls_profile(path))
        elif suffix == ".shp":
            profile.update(status="schema_profiled", **shapefile_profile(path))
        elif suffix == ".dbf":
            profile.update(status="schema_profiled", **dbf_profile(path))
        elif suffix == ".zip":
            with zipfile.ZipFile(path) as archive:
                members = archive.infolist()
                profile.update(
                    status="archive_profiled",
                    members=len(members),
                    member_names=[member.filename for member in members],
                    uncompressed_bytes=sum(member.file_size for member in members),
                )
        elif suffix == ".pdf":
            from pypdf import PdfReader

            pdf = PdfReader(str(path))
            profile.update(
                status="document_profiled",
                pages=len(pdf.pages),
                encrypted=bool(pdf.is_encrypted),
            )
        elif suffix == ".docx":
            from docx import Document

            document = Document(str(path))
            profile.update(
                status="document_profiled",
                paragraphs=len(document.paragraphs),
                tables=len(document.tables),
            )
        elif suffix == ".pptx":
            from pptx import Presentation

            deck = Presentation(str(path))
            profile.update(status="document_profiled", slides=len(deck.slides))
        elif suffix in {
            ".md",
            ".txt",
            ".log",
            ".bib",
            ".ini",
            ".html",
            ".prj",
            ".cpg",
            ".cst",
        }:
            text = path.read_text(encoding="utf-8", errors="replace")
            profile.update(
                status="text_profiled", characters=len(text), lines=len(text.splitlines())
            )
    except Exception as exc:  # Preserve coverage while making failures visible.
        profile.update(status="profile_error", error=f"{type(exc).__name__}: {exc}")
    return profile


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("root", type=Path)
    parser.add_argument("output_dir", type=Path)
    args = parser.parse_args()

    root = args.root.resolve()
    output_dir = args.output_dir.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    files = sorted(
        path
        for path in root.rglob("*")
        if path.is_file() and "graphify-out" not in path.relative_to(root).parts
    )
    profiles = []
    for path in files:
        stat = path.stat()
        relative = path.relative_to(root).as_posix()
        profiles.append(
            {
                "relative_path": relative,
                "extension": path.suffix.lower() or "[none]",
                "size_bytes": stat.st_size,
                "modified_utc": datetime.fromtimestamp(
                    stat.st_mtime, tz=timezone.utc
                ).isoformat(),
                "sha256": sha256(path),
                "profile": basic_profile(path),
            }
        )

    generated_at = datetime.now(timezone.utc).isoformat()
    payload = {
        "snapshot_generated_utc": generated_at,
        "source_root": str(root),
        "source_mutated": False,
        "excluded_generated_directories": ["graphify-out"],
        "file_count": len(profiles),
        "total_bytes": sum(item["size_bytes"] for item in profiles),
        "files": profiles,
    }
    (output_dir / "dropbox_corpus_profiles_2026-08-31.json").write_text(
        json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8"
    )

    with (output_dir / "dropbox_corpus_manifest_2026-08-31.csv").open(
        "w", newline="", encoding="utf-8-sig"
    ) as stream:
        writer = csv.DictWriter(
            stream,
            fieldnames=[
                "relative_path",
                "extension",
                "size_bytes",
                "modified_utc",
                "sha256",
                "profile_status",
            ],
        )
        writer.writeheader()
        for item in profiles:
            writer.writerow(
                {
                    **{key: item[key] for key in writer.fieldnames[:-1]},
                    "profile_status": item["profile"]["status"],
                }
            )

    ext_counts = Counter(item["extension"] for item in profiles)
    top_counts = Counter(item["relative_path"].split("/", 1)[0] for item in profiles)
    status_counts = Counter(item["profile"]["status"] for item in profiles)
    errors = [item for item in profiles if item["profile"]["status"] == "profile_error"]

    lines = [
        "# Dropbox corpus snapshot — 31 August 2026",
        "",
        "This is a read-only, content-aware inventory of the Advancing Justice Uganda project Dropbox. "
        "Graphify-generated conversion sidecars are excluded from the source count.",
        "",
        "## Coverage",
        "",
        f"- User-owned files: **{len(profiles)}**",
        f"- Total size: **{sum(item['size_bytes'] for item in profiles):,} bytes**",
        "- Every file has a SHA-256 checksum and modification timestamp.",
        "- Tabular and geospatial datasets receive schema profiles; archives receive member inventories; documents receive structural profiles.",
        "- Source Dropbox files were not modified by this snapshot script.",
        "",
        "### Files by top-level area",
        "",
    ]
    lines += [f"- `{name}`: {count}" for name, count in sorted(top_counts.items())]
    lines += ["", "### Files by extension", ""]
    lines += [f"- `{ext}`: {count}" for ext, count in ext_counts.most_common()]
    lines += ["", "### Processing status", ""]
    lines += [f"- `{status}`: {count}" for status, count in status_counts.most_common()]
    lines += ["", "## Canonical August 2026 analytical context", ""]
    lines += [
        "- `BASELINE_RECONCILIATION_AND_ANALYTICAL_GUIDANCE_AUG2026.md` is the controlling project note for the revised Phase 1 baseline workflow.",
        "- The earlier 133-record analysis is a pre-amendment release and must not be treated as the final baseline.",
        "- The final cohort is a lineage problem, not an append/drop operation: reconcile villages, respondents, LC1 election turnover, mop-up surveys, replacements, and actual training exposure.",
        "- Preserve every raw submission, then select exactly one valid pre-treatment record for each final intervention LC/LCC.",
        "- The administrative election crosswalk is `3 Data/1 Raw/Secondary data/List of LCs mapped.xlsx`.",
        "- The original survey is `3 Data/1 Raw/Primary data/Phase 1 Baseline_Runyankore_WIDE.xlsx`; the August mop-up file is `Phase 1 Baseline_Runyankore_WIDE (new submissions).csv` with 34 submissions.",
        "- Thirteen rows are marked `VOTED OUT`; at least eight have obvious post-election re-baselines. Supersede respondent records where appropriate rather than automatically deleting villages.",
        "- Release remains blocked until actual training dates prove each retained survey occurred before first treatment exposure.",
        "",
        "## Required data architecture",
        "",
        "1. Archival all-submission raw file.",
        "2. Canonical village crosswalk.",
        "3. Record-level lineage and supersession ledger.",
        "4. Harmonized clean file retaining all records.",
        "5. Final analysis file with one verified pre-treatment record per final LC/LCC.",
        "6. Reconciliation, QA, release-status, and sample-flow outputs generated from code.",
        "",
        "## Profile exceptions",
        "",
    ]
    if errors:
        lines += [
            f"- `{item['relative_path']}` — {item['profile'].get('error')}" for item in errors
        ]
    else:
        lines.append("- None.")
    lines += [
        "",
        "## Machine-readable companions",
        "",
        "- `dropbox_corpus_manifest_2026-08-31.csv`: one row per source file.",
        "- `dropbox_corpus_profiles_2026-08-31.json`: detailed schemas, archive members, structural profiles, hashes, and timestamps.",
    ]
    (output_dir / "DROPBOX_CONTEXT_2026-08-31.md").write_text(
        "\n".join(lines) + "\n", encoding="utf-8"
    )
    print(json.dumps({"files": len(profiles), "statuses": status_counts, "errors": len(errors)}, default=dict))


if __name__ == "__main__":
    main()
